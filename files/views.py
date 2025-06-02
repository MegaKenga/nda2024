from django.shortcuts import render, redirect, get_object_or_404
from openpyxl import load_workbook

from catalog.models import Offer, Product


def import_from_excel(request, product_id):
    if request.method == 'POST':
        try:
            excel_file = request.FILES['excel_file']
            wb = load_workbook(excel_file)
            ws = wb.active

            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    name, text_description, shipping_pack, place = row
                    product_id = int(product_id)
                    Offer.objects.create(
                        name=name,
                        text_description=text_description,
                        shipping_pack=shipping_pack,
                        place=place,
                        status='PUBLISHED',
                        product_id=product_id
                    )
                else:
                    break
            product = get_object_or_404(Product, pk=int(product_id))
            return redirect('offer', product_slug=product.slug, brand_slug=product.brand.slug)
        except Exception:
            product = get_object_or_404(Product, pk=int(product_id))
            return redirect('offer', product_slug=product.slug, brand_slug=product.brand.slug)

    return render(request, 'core/components/excel_input.html')


def delete_offers(request, product_id):
    if request.method == 'POST':
        try:
            get_offers = Offer.objects.filter(product=product_id)
            get_offers.delete()
            product = get_object_or_404(Product, pk=int(product_id))
            return redirect('offer', product_slug=product.slug, brand_slug=product.brand.slug)
        except Exception:
            product = get_object_or_404(Product, pk=int(product_id))
            return redirect('offer', product_slug=product.slug, brand_slug=product.brand.slug)

    return render(request, 'core/components/delete_offers.html')